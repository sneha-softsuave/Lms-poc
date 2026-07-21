"""Phase 2 smoke test: AI course generation + review/publish.

Uses a FakeGateway (canned JSON) injected via FastAPI dependency_overrides, so
it runs offline with no API key. Exercises: ingest -> generate draft course ->
inspect tree + provenance + pending questions -> publish blocked while pending
-> approve/reject -> publish (version bump).

Run: .venv/Scripts/python.exe scripts/smoke_phase2.py
"""

import asyncio
import io
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

_tmp = tempfile.mktemp(suffix=".db")
os.environ["DATABASE_URL"] = f"sqlite+aiosqlite:///{_tmp}"
os.environ["STORAGE_PROVIDER"] = "local"
os.environ["LOCAL_STORAGE_PATH"] = tempfile.mkdtemp()

import httpx  # noqa: E402

from app.gateway.base import Completion  # noqa: E402
from app.gateway.factory import get_gateway  # noqa: E402
from app.main import app  # noqa: E402

COURSE_JSON = """{
  "subject": {"title": "Vehicle Mechanics", "description": "Assemblies and first-line checks"},
  "course": {
    "title": "Vehicle Mechanics - Fundamentals",
    "description": "Powerpack basics",
    "objectives": ["Identify major assemblies", "Perform first-line checks"],
    "modules": [
      {"title": "Powerpack", "lessons": [
        {"title": "Engine overview", "body": "The powerpack is the engine assembly.",
         "source_ref": {"doc": "VEH-01", "section": "Engine Overview", "page": 1}},
        {"title": "Cooling system", "body": "The coolant pump prevents overheating.",
         "source_ref": {"doc": "VEH-01", "section": "Cooling System", "page": 2}}
      ]}
    ],
    "glossary": [
      {"term": "Coolant pump", "definition": "Circulates fluid to cool the engine.",
       "source_ref": {"doc": "VEH-01", "section": "Cooling System", "page": 2}}
    ]
  }
}"""

QUIZ_JSON = """{
  "questions": [
    {"qtype": "mcq", "stem": "What does the coolant pump do?",
     "options": ["Cools the engine", "Ignites fuel", "Steers", "Brakes"],
     "correct_answer": "Cools the engine", "rationale": "It circulates coolant.",
     "difficulty": "basic", "source_ref": {"doc": "VEH-01", "section": "Cooling System", "page": 2}},
    {"qtype": "tf", "stem": "The powerpack is the engine assembly.",
     "options": ["true", "false"], "correct_answer": "true", "rationale": "Stated in the material.",
     "difficulty": "basic", "source_ref": {"doc": "VEH-01", "section": "Engine Overview", "page": 1}}
  ]
}"""

QUALITY_JSON = '{"overall_score": 88, "clarity": 22, "correctness": 24, "difficulty": 18, "format": 14, "educational": 10}'


class FakeGateway:
    """Canned responses keyed by which system prompt is used."""

    embedding_dim = 384

    async def generate(self, prompt, *, system=None, max_tokens=1024, temperature=0.0, stop=None):
        s = system or ""
        if "curriculum architect" in s:
            return Completion(text=COURSE_JSON)
        if "assessment questions" in s:
            return Completion(text=QUIZ_JSON)
        if "assessment reviewer" in s:
            return Completion(text=QUALITY_JSON)
        return Completion(text="{}")

    async def embed(self, texts):
        return [[0.0] * 384 for _ in texts]

    async def rerank(self, query, passages):
        return [0.0 for _ in passages]


def make_docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Engine Overview", level=1)
    d.add_paragraph("The powerpack is the vehicle's engine assembly.")
    d.add_heading("Cooling System", level=1)
    d.add_paragraph("The coolant pump circulates fluid to prevent overheating at idle.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


async def main() -> None:
    app.dependency_overrides[get_gateway] = lambda: FakeGateway()
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as c:
        async with app.router.lifespan_context(app):
            token = (await c.post("/api/v1/auth/login",
                     json={"email": "admin@defense-lms.org", "password": "admin12345"})).json()["access_token"]
            H = {"Authorization": f"Bearer {token}"}

            # ingest
            doc = (await c.post("/api/v1/ingestion/upload", headers=H,
                   files={"file": ("veh.docx", make_docx(),
                          "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})).json()

            # generate draft course
            r = await c.post(f"/api/v1/course-generation/documents/{doc['id']}/generate", headers=H)
            assert r.status_code == 201, r.text
            course = r.json()
            assert course["status"] == "draft", course["status"]
            assert course["title"].startswith("Vehicle Mechanics"), course["title"]
            assert len(course["modules"]) == 1 and len(course["modules"][0]["lessons"]) == 2
            # provenance propagated to lessons
            assert course["modules"][0]["lessons"][1]["source_ref"]["page"] == 2
            assert len(course["glossary"]) == 1
            course_id = course["id"]
            print("generate: draft course with", len(course["modules"]), "module(s), provenance ok")

            # questions generated as pending, with quality scores
            qs = (await c.get(f"/api/v1/review/courses/{course_id}/questions", headers=H)).json()
            assert len(qs) == 2, qs
            assert all(q["review_status"] == "pending" for q in qs)
            assert all(q["quality_score"] == 88 for q in qs)
            print("quiz bank:", len(qs), "pending questions, quality scored")

            # publish blocked while pending
            r = await c.post(f"/api/v1/review/courses/{course_id}/publish", headers=H)
            assert r.status_code == 409, f"expected 409, got {r.status_code}: {r.text}"
            print("publish blocked while pending: ok")

            # approve one, reject one, edit one
            await c.post(f"/api/v1/review/questions/{qs[0]['id']}/approve", headers=H)
            await c.post(f"/api/v1/review/questions/{qs[1]['id']}/reject", headers=H)
            r = await c.patch(f"/api/v1/review/questions/{qs[0]['id']}", headers=H,
                              json={"difficulty": "intermediate"})
            assert r.json()["difficulty"] == "intermediate"

            # attach a 3D model to a lesson
            les_id = course["modules"][0]["lessons"][0]["id"]
            r = await c.patch(f"/api/v1/review/lessons/{les_id}", headers=H,
                              json={"model3d_id": "engine-cutaway"})
            assert r.json()["model3d_id"] == "engine-cutaway"

            # now publish succeeds + version bump
            r = await c.post(f"/api/v1/review/courses/{course_id}/publish", headers=H)
            assert r.status_code == 200, r.text
            assert r.json()["status"] == "published" and r.json()["version"] == 2, r.json()
            print("publish after review: ok, version", r.json()["version"])

    app.dependency_overrides.clear()
    print("\nPHASE 2 SMOKE: PASS")


if __name__ == "__main__":
    asyncio.run(main())
